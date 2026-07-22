"""Format adapters producing one normalized content document."""

from __future__ import annotations

import csv
import io
import json
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from domain.content.models import (
    Answer,
    CompetencyMapping,
    ContentDocument,
    Difficulty,
    Domain,
    Exercise,
    Explanation,
    Hint,
    Media,
    Prerequisite,
    Program,
    Question,
    Skill,
    Subject,
    SubSkill,
    Tag,
    ValidationStatus,
    Version,
)


class ContentImportError(ValueError):
    """Raised when an input cannot be normalized safely."""


class ContentImporter(ABC):
    @abstractmethod
    def load(self, source: str | Path) -> ContentDocument: ...


def _mapping(source: str | Path, parser: Any) -> dict[str, Any]:
    text = Path(source).read_text(encoding="utf-8") if isinstance(source, Path) else source
    value = parser(text)
    if not isinstance(value, dict):
        raise ContentImportError("The content root must be an object")
    return value


def _version(data: dict[str, Any] | None) -> Version:
    data = data or {}
    return Version(
        int(data.get("number", 1)), str(data.get("author", "unknown")), ValidationStatus(data.get("status", "draft"))
    )


def document_from_mapping(data: dict[str, Any]) -> ContentDocument:
    subjects = tuple(Subject(str(x["code"]), str(x["label"])) for x in data.get("subjects", []))
    programs = tuple(
        Program(str(x["code"]), str(x["label"]), str(x["country_code"]), _version(x.get("version")))
        for x in data.get("programs", [])
    )
    domains = tuple(Domain(str(x["code"]), str(x["label"]), str(x["subject_code"])) for x in data.get("domains", []))
    skills = tuple(
        Skill(
            str(x["code"]),
            str(x["label"]),
            str(x["domain_code"]),
            tuple(Prerequisite(str(x["code"]), str(p)) for p in x.get("prerequisites", [])),
        )
        for x in data.get("skills", [])
    )
    subskills = tuple(
        SubSkill(str(x["code"]), str(x["label"]), str(x["skill_code"])) for x in data.get("subskills", [])
    )
    exercises: list[Exercise] = []
    for item in data.get("exercises", []):
        questions: list[Question] = []
        for raw in item.get("questions", []):
            questions.append(
                Question(
                    code=str(raw["code"]),
                    statement=str(raw.get("statement", "")),
                    answer=Answer(raw.get("answer"), str(raw.get("answer_type", "text"))),
                    explanation=Explanation(str(raw.get("explanation", ""))),
                    difficulty=Difficulty(int(raw.get("difficulty", item.get("difficulty", 3)))),
                    mappings=tuple(
                        CompetencyMapping(
                            str(m["skill_code"]), float(m.get("weight", 1)), bool(m.get("primary", False))
                        )
                        for m in raw.get("skills", [])
                    ),
                    hints=tuple(Hint(str(h), position) for position, h in enumerate(raw.get("hints", []), 1)),
                    tags=tuple(Tag(str(t), str(t)) for t in raw.get("tags", [])),
                    media=tuple(
                        Media(
                            str(m["code"]), str(m["kind"]), str(m["uri"]), str(m.get("title", "")), m.get("mime_type")
                        )
                        for m in raw.get("media", [])
                    ),
                    version=_version(raw.get("version")),
                )
            )
        exercises.append(
            Exercise(
                str(item["code"]),
                str(item["title"]),
                str(item.get("objective", "")),
                str(item["subject_code"]),
                Difficulty(int(item.get("difficulty", 3))),
                tuple(questions),
                tuple(Tag(str(t), str(t)) for t in item.get("tags", [])),
                tuple(
                    Media(str(m["code"]), str(m["kind"]), str(m["uri"]), str(m.get("title", "")), m.get("mime_type"))
                    for m in item.get("media", [])
                ),
                _version(item.get("version")),
            )
        )
    return ContentDocument(
        programs, subjects, domains, skills, subskills, tuple(exercises), dict(data.get("metadata", {}))
    )


class JSONContentImporter(ContentImporter):
    def load(self, source: str | Path) -> ContentDocument:
        return document_from_mapping(_mapping(source, json.loads))


class YAMLContentImporter(ContentImporter):
    def load(self, source: str | Path) -> ContentDocument:
        try:
            import yaml  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ContentImportError("YAML import requires PyYAML") from exc
        return document_from_mapping(_mapping(source, yaml.safe_load))


class CSVContentImporter(ContentImporter):
    """Import normalized rows with an ``entity_type`` discriminator."""

    def load(self, source: str | Path) -> ContentDocument:
        text = source.read_text(encoding="utf-8") if isinstance(source, Path) else source
        grouped: dict[str, list[dict[str, Any]]] = {}
        for row in csv.DictReader(io.StringIO(text)):
            kind = row.pop("entity_type", "").strip()
            if not kind:
                raise ContentImportError("CSV rows require entity_type")
            cleaned = {key: value for key, value in row.items() if value not in (None, "")}
            grouped.setdefault(kind, []).append(cleaned)
        return document_from_mapping(grouped)


class MarkdownContentImporter(ContentImporter):
    """Import a Markdown document containing a fenced JSON manifest."""

    def load(self, source: str | Path) -> ContentDocument:
        text = source.read_text(encoding="utf-8") if isinstance(source, Path) else source
        marker = "```json"
        if marker not in text:
            raise ContentImportError("Markdown content requires a fenced JSON manifest")
        manifest = text.split(marker, 1)[1].split("```", 1)[0]
        return JSONContentImporter().load(manifest)


class ExcelContentImporter(ContentImporter):
    """Import sheets named after plural entity collections."""

    def load(self, source: str | Path) -> ContentDocument:
        if not isinstance(source, Path):
            raise ContentImportError("Excel import requires a file path")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ContentImportError("Excel import requires openpyxl") from exc
        workbook = load_workbook(source, read_only=True, data_only=True)
        data: dict[str, list[dict[str, Any]]] = {}
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            data[sheet.title] = [dict(zip((str(h) for h in headers), row, strict=True)) for row in rows]
        return document_from_mapping(data)


class ImporterRegistry:
    def __init__(self) -> None:
        self._importers: dict[str, ContentImporter] = {}

    def register(self, suffix: str, importer: ContentImporter) -> None:
        self._importers[suffix.lower().lstrip(".")] = importer

    def for_source(self, source: Path) -> ContentImporter:
        suffix = source.suffix.lower().lstrip(".")
        try:
            return self._importers[suffix]
        except KeyError as exc:
            raise ContentImportError(f"Unsupported content format: {suffix or '<none>'}") from exc

    @classmethod
    def defaults(cls) -> ImporterRegistry:
        registry = cls()
        for suffix, importer in {
            "json": JSONContentImporter(),
            "csv": CSVContentImporter(),
            "md": MarkdownContentImporter(),
            "markdown": MarkdownContentImporter(),
            "yaml": YAMLContentImporter(),
            "yml": YAMLContentImporter(),
            "xlsx": ExcelContentImporter(),
        }.items():
            registry.register(suffix, importer)
        return registry
